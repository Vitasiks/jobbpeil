import csv
import io
import math
import re
import sys
import warnings
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
from html.parser import HTMLParser
from http.client import HTTPException
from urllib.error import URLError
from urllib.parse import urljoin, urlsplit
from urllib.request import urlopen


NAV_PAGE = (
    "https://www.nav.no/no/nav-og-samfunn/statistikk/"
    "arbeidssokere-og-stillinger-statistikk/hovedtall-om-arbeidsmarkedet"
)
COUNTIES = ("Agder", "Rogaland", "Vestland")
VACANCIES_PAGE = NAV_PAGE.rsplit("/", 1)[0] + "/ledige-stillinger"
MONTHS = (
    "januar", "februar", "mars", "april", "mai", "juni",
    "juli", "august", "september", "oktober", "november", "desember",
)


class LinkParser(HTMLParser):
    """Собирает ссылки вместе с их видимым текстом."""

    def __init__(self):
        super().__init__()
        self.links = []
        self.href = None
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.href = dict(attrs).get("href")
            self.parts = []

    def handle_data(self, data):
        if self.href is not None:
            self.parts.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.href is not None:
            self.links.append((self.href, "".join(self.parts)))
            self.href = None


def download(url):
    with urlopen(url, timeout=30) as response:
        return response.read()


def find_latest_csv(html):
    parser = LinkParser()
    parser.feed(html)
    candidates = []
    for href, label in parser.links:
        label = " ".join(label.lower().split())
        url = urljoin(NAV_PAGE, href)
        if "helt ledige etter fylke og kommune" not in label:
            continue
        if not urlsplit(url).path.lower().endswith(".csv"):
            continue
        match = re.search(r"\b(" + "|".join(MONTHS) + r")\s+(\d{4})\b", label)
        if match:
            month, year = match.groups()
            candidates.append((int(year), MONTHS.index(month) + 1, url))

    if not candidates:
        raise ValueError("На странице NAV не найдена датированная ссылка на нужный CSV.")
    year, month, url = max(candidates)
    return url, f"{month:02d}.{year}"


def read_counties(raw):
    # NAV публикует CSV с разделителем ';'; возможны UTF-8 и Windows-1252.
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")

    reader = csv.DictReader(io.StringIO(text), delimiter=";", strict=True)
    required = {"Fylkenr", "Fylkenavn", "Kommunenr", "Antall helt ledige"}
    if not required.issubset(reader.fieldnames or []):
        raise ValueError("В CSV отсутствуют необходимые столбцы.")

    percent_column = "Helt ledige i prosent av arbeidsstyrken"
    results = {}
    for row in reader:
        if None in row or any(value is None for value in row.values()):
            raise ValueError("В CSV обнаружена строка неправильной длины.")
        name = row["Fylkenavn"].strip()
        # В итоговой строке fylke код Kommunenr совпадает с Fylkenr.
        if name not in COUNTIES or row["Kommunenr"].strip() != row["Fylkenr"].strip():
            continue
        if name in results:
            raise ValueError(f"В CSV повторяется итоговая строка {name}.")
        count = int(row["Antall helt ledige"].strip())
        if count < 0:
            raise ValueError(f"Некорректное количество безработных для {name}.")
        percent = None
        if percent_column in row:
            percent = float(row[percent_column].strip().replace(",", "."))
            if not math.isfinite(percent) or not 0 <= percent <= 100:
                raise ValueError(f"Некорректный процент для {name}.")
        results[name] = (count, percent)

    if results.keys() != set(COUNTIES):
        raise ValueError("В CSV найдены итоговые данные не для всех трёх fylke.")
    return results


def find_vacancies_xlsx(html, period, category="fylke"):
    """Находит публикацию вакансий за месяц выбранного CSV безработицы."""
    month, year = period.split(".")
    expected_date = f"{MONTHS[int(month) - 1]} {year}"
    parser = LinkParser()
    parser.feed(html)
    for href, label in parser.links:
        label = " ".join(label.lower().split())
        url = urljoin(VACANCIES_PAGE, href)
        if (
            f"tilgang ledige stillinger. {category}." in label
            and re.search(r"\b" + re.escape(expected_date) + r"\b", label)
            and urlsplit(url).path.lower().endswith(".xlsx")
        ):
            return url
    raise ValueError(f"На странице NAV нет XLSX вакансий ({category}) за {period}.")


def read_vacancy_rows(raw, period, sheet_name):
    """Читает абсолютное месячное количество, проверяя год и заголовок месяца."""
    # Учебный numbers.py рядом с программой перекрывает стандартный модуль.
    # Исключаем каталог проекта только на время импорта библиотеки.
    original_path = sys.path[:]
    try:
        project_dir = Path(__file__).resolve().parent
        sys.path[:] = [p for p in sys.path if Path(p).resolve() != project_dir]
        from openpyxl import load_workbook
    finally:
        sys.path[:] = original_path

    month, year = period.split(".")
    month_name = MONTHS[int(month) - 1]
    with warnings.catch_warnings():
        # В файле NAV нет стиля по умолчанию; на значения ячеек это не влияет.
        warnings.filterwarnings(
            "ignore", message="Workbook contains no default style.*", category=UserWarning
        )
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"В XLSX нет листа {sheet_name}.")
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        period_headers = [
            str(cell).strip().lower()
            for row in rows for cell in row
            if isinstance(cell, str)
            and re.fullmatch(r"januar\s*[-–]\s*\w+\s+\d{4}", cell.strip().lower())
        ]
        if period_headers != [f"januar - {month_name} {year}"]:
            raise ValueError("Период внутри XLSX не совпадает с периодом безработицы.")
        columns = {
            index for row in rows for index, cell in enumerate(row)
            if isinstance(cell, str) and cell.strip().lower() == month_name
        }
        if len(columns) != 1:
            raise ValueError("В XLSX не найден однозначный столбец нужного месяца.")
        return rows, columns.pop()
    finally:
        workbook.close()


def vacancy_count(value, name):
    if (
        type(value) not in (int, float) or not math.isfinite(value)
        or value < 0 or value != int(value)
    ):
        raise ValueError(f"В XLSX некорректное количество вакансий для {name}.")
    return int(value)


def read_vacancies(raw, period):
    rows, column = read_vacancy_rows(raw, period, "1. Antall. Fylke")
    results = {}
    for row in rows:
        names = [cell.strip() for cell in row if isinstance(cell, str) and cell.strip() in COUNTIES]
        if not names:
            continue
        if len(names) != 1 or names[0] in results:
            raise ValueError("В XLSX неоднозначные строки fylke.")
        results[names[0]] = vacancy_count(row[column], names[0])
    if results.keys() != set(COUNTIES):
        raise ValueError("В XLSX нет вакансий для всех трёх fylke.")
    return results


def read_occupation_vacancies(raw, period):
    rows, column = read_vacancy_rows(raw, period, "1. Antall")
    results = {}
    # Каждый раздел начинается названием группы и строкой заголовков месяцев.
    # Берём строку «I alt ...», чтобы не смешивать итоги и подгруппы.
    month_name = MONTHS[int(period.split(".")[0]) - 1]
    group = None
    for index, row in enumerate(rows):
        if isinstance(row[column], str) and row[column].strip().lower() == month_name:
            if group is not None or index == 0:
                raise ValueError("В XLSX нет итога предыдущей профессиональной группы.")
            heading = rows[index - 1][1]
            if not isinstance(heading, str) or not heading.strip():
                raise ValueError("В XLSX отсутствует название профессиональной группы.")
            group = heading.strip()
            continue
        label = row[1]
        if isinstance(label, str) and label.strip().lower().startswith("i alt "):
            if group is None or label.strip().lower() != f"i alt {group.lower()}":
                raise ValueError("Итог XLSX не соответствует профессиональной группе.")
            if group in results:
                raise ValueError(f"В XLSX повторяется профессиональная группа {group}.")
            results[group] = vacancy_count(row[column], group)
            group = None
    if group is not None or not results:
        raise ValueError("В XLSX нет полных итогов профессиональных групп.")
    return results


def main():
    try:
        html = download(NAV_PAGE).decode("utf-8-sig")
        csv_url, period = find_latest_csv(html)
        results = read_counties(download(csv_url))
        vacancies_html = download(VACANCIES_PAGE).decode("utf-8-sig")
        vacancies_url = find_vacancies_xlsx(vacancies_html, period)
        vacancies = read_vacancies(download(vacancies_url), period)
        occupations_url = find_vacancies_xlsx(vacancies_html, period, "yrke")
        occupations = read_occupation_vacancies(download(occupations_url), period)
    except ImportError:
        print("Ошибка: для чтения XLSX необходима библиотека openpyxl.")
        return
    except (URLError, OSError, HTTPException):
        print("Ошибка: не удалось загрузить данные NAV.")
        return
    except (ValueError, csv.Error, BadZipFile, ParseError, KeyError, IndexError) as error:
        print(f"Ошибка чтения данных NAV: {error}")
        return

    print(f"Период: {period}")
    ratios = {}
    for name in COUNTIES:
        count, percent = results[name]
        line = f"{name}: полностью безработных — {count}"
        if percent is not None:
            line += f"; от рабочей силы — {percent:g}%"
        line += f"; новых вакансий — {vacancies[name]}"
        if count > 0:
            ratios[name] = vacancies[name] / count * 100
            line += f"; новых вакансий на 100 безработных — {ratios[name]:.2f}"
        else:
            line += "; новых вакансий на 100 безработных — не рассчитано (безработных 0)"
        print(line)

    if ratios:
        best_region = max(ratios, key=ratios.get)
        print(
            f"Самый высокий показатель: {best_region} — "
            f"{ratios[best_region]:.2f} новых вакансий на 100 полностью безработных."
        )
    else:
        print("Регион с самым высоким показателем не определён: безработных 0 во всех регионах.")

    print(f"\nНовые вакансии по профессиональным группам — вся Норвегия, {period}:")
    for name, count in occupations.items():
        print(f"{name}: {count}")


if __name__ == "__main__":
    main()
